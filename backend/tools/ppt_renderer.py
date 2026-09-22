# 目录: ./backend/tools/ppt_renderer.py | 模块: PPT 渲染 | 职责: 套用 templates/ 下模板的真实设计（封面/目录/致谢页复用）并渲染每页内容
import os
import re
import copy
import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

from core.state import AgentState
from utils.progress import update_progress, push_event, log_step

TEMPLATE_DIR = "./templates"
DEFAULT_TEMPLATE = "master.pptx"
OUTPUT_DIR = "./output"

# 取自模板红色色带（矩形填充 9B0000），保证新增文字与模板视觉一致
ACCENT = RGBColor(0x9B, 0x00, 0x00)
SUBTITLE_COLOR = RGBColor(0x40, 0x40, 0x40)
TITLE_FONT = "微软雅黑"
TOC_PATTERN = re.compile(r"^\s*(\d+)\s*[、.．]")
# LLM 有时会把"封面""公司logo"这类设计说明当成封面副标题写进 bullets，展示时过滤掉
META_HINT = re.compile(r"封面|公司\s*logo|logo|汇报人|目录页|致谢页|占位|设计说明|背景图|报告标题", re.I)

# LLM/planner 可能给出英文版式名，这里兜底映射到模板里的真实中文版式名
LAYOUT_ALIASES = {
    "title slide": "标题幻灯片",
    "title and content": "标题和内容",
    "two content": "两栏内容",
    "title only": "仅标题",
    "section header": "节标题",
    "blank": "空白",
}


def resolve_layout(prs: Presentation, layout_name: str):
    """按 名称精确匹配 → 别名映射 → 结构匹配 的顺序找版式，避免全部退化成第 0 个版式。"""
    layouts = list(prs.slide_layouts)
    for layout in layouts:
        if layout.name == layout_name:
            return layout

    alias = LAYOUT_ALIASES.get(layout_name.strip().lower())
    if alias:
        for layout in layouts:
            if layout.name == alias:
                return layout

    wants_title_only = alias == "仅标题" or layout_name == "仅标题"
    for layout in layouts:
        idxs = {p.placeholder_format.idx for p in layout.placeholders}
        has_title = 0 in idxs
        has_body = 1 in idxs
        if has_title and (has_body != wants_title_only):
            return layout
    return layouts[0]


def pick_template() -> str:
    """模板选择：优先 PPT_TEMPLATE 环境变量，其次 master.pptx，最后取目录下任意 .pptx"""
    name = os.getenv("PPT_TEMPLATE") or DEFAULT_TEMPLATE
    path = os.path.join(TEMPLATE_DIR, name)
    if os.path.exists(path):
        return path
    if os.path.isdir(TEMPLATE_DIR):
        candidates = sorted(f for f in os.listdir(TEMPLATE_DIR) if f.lower().endswith(".pptx"))
        if candidates:
            return os.path.join(TEMPLATE_DIR, candidates[0])
    raise FileNotFoundError(f"模板目录 {TEMPLATE_DIR} 下没有可用的 .pptx 模板")


def slide_text(slide) -> str:
    return "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def set_paragraph_text(paragraph, text: str):
    """写入文字并尽量保留模板原有的 run 级字体/字号/颜色"""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = text


def add_textbox(slide, lines, left, top, width, height, size, *, bold=False,
                color=SUBTITLE_COLOR, align=PP_ALIGN.CENTER):
    """在复用的模板页上新增文本框（模板的封面/致谢页本身没有标题占位符）"""
    if isinstance(lines, str):
        lines = [lines]
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size if i == 0 else max(size - 6, 12))
        run.font.bold = bold if i == 0 else False
        run.font.color.rgb = color
        run.font.name = TITLE_FONT
        # 中文字形走 ea/cs，否则 PowerPoint 会用 latin 字体渲染中文
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                rPr.append(el)
            el.set("typeface", TITLE_FONT)
    return box


def fill_date_chrome(slide, today: str):
    """模板里的『汇报时间：』空值自动补上生成日期"""
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        text = sh.text_frame.text.strip()
        if text.startswith("汇报时间") and len(text) <= 12:
            set_paragraph_text(sh.text_frame.paragraphs[0], f"{text}{today}")


def _toc_entries(slide) -> list:
    """收集目录页的条目框，返回 [(编号, 文档序号, shape)]，按视觉顺序（编号）排列"""
    entries = []
    for i, sh in enumerate(slide.shapes):
        if not sh.has_text_frame:
            continue
        first = sh.text_frame.paragraphs[0].text if sh.text_frame.paragraphs else ""
        if re.match(r"^\s*目录", first):
            continue
        m = TOC_PATTERN.match(first or "")
        if m:
            entries.append((int(m.group(1)), i, sh))
    entries.sort(key=lambda e: (e[0], e[1]))
    return entries


def fill_toc(slide, titles) -> bool:
    """目录页：把 1、2、3… 条目替换成真实页标题。

    模板默认只有 5 个条目框（微软雅黑/28pt/C00000）。当生成的目录条目更多时：
    1. 深拷贝克隆条目框 —— 继承模板 run 级字体/字号/颜色，避免退回默认黑色；
    2. 在模板原有的纵向范围内等距重排，铺满整页；
    3. 按"条目数、标题长度、行距"自动缩小字号（下限 12pt）。
    条目数少于模板时删掉多余的空框。
    """
    entries = _toc_entries(slide)
    if not entries:
        return False

    shapes = [e[2] for e in entries]
    capacity = len(shapes)
    n = len(titles)

    if n > capacity:
        # 克隆最后一个条目框补足数量（deepcopy 保留 spPr/rPr 即模板样式）
        src_el = shapes[-1]._element
        for _ in range(n - capacity):
            slide.shapes._spTree.append(copy.deepcopy(src_el))
        entries = _toc_entries(slide)  # 克隆体编号与末框相同，按(编号, 文档序)排即正确
        shapes = [e[2] for e in entries]
    elif n < capacity:
        for sh in shapes[n:]:
            sh._element.getparent().remove(sh._element)
        shapes = shapes[:n]

    # 基准几何：模板原有的条目位置/尺寸/步距
    base = shapes[:capacity] if n > capacity else shapes
    tops = [sh.top for sh in base]
    y_min = min(tops)
    box_h = base[0].height
    y_max = max(tops) + box_h
    span = y_max - y_min - box_h
    template_step = span / (capacity - 1) if capacity > 1 else box_h

    # 条目变多：在原有纵向范围内等距铺开
    if n > capacity and n > 1:
        step = span / (n - 1)
        for idx, sh in enumerate(shapes):
            sh.top = int(round(y_min + idx * step))

    # 自动缩字号：保证两行内放得下、两行不超出步距、条目变多按比例缩
    from pptx.util import Emu  # noqa: WPS433  函数内导入，避免顶层依赖
    orig_size = None
    for sh in base:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    orig_size = r.font.size.pt
                    break
            if orig_size:
                break
        if orig_size:
            break
    orig_size = orig_size or 28.0  # 模板默认 28pt 兜底

    effective_step = template_step if n <= capacity else (span / (n - 1) if n > 1 else box_h)
    box_w_in = Emu(base[0].width).inches
    max_len = max((len(str(t)) for t in titles), default=1)

    size = orig_size
    if max_len > 0:
        size = min(size, 2 * box_w_in * 72 / max_len)     # 最多两行放下
    size = min(size, effective_step * 0.95 * 72 / 2)      # 两行不超出步距
    if n > capacity:
        size = min(size, orig_size * capacity / n)        # 条目变多按比例缩
    size = max(12, int(round(size)))

    # 写入标题（保持编号）并统一下发缩放后的字号
    for i, sh in enumerate(shapes):
        p = sh.text_frame.paragraphs[0]
        set_paragraph_text(p, f"{i + 1}、{titles[i]}" if i < len(titles) else "")
        if size < orig_size:
            for run in p.runs:
                run.font.size = Pt(size)
        if not (p.text or "").strip():
            sh._element.getparent().remove(sh._element)  # 多余的空框直接删掉
    return True


def subtitle_lines(slide_data: dict, limit: int) -> list:
    """封面/致谢页副标题：过滤掉设计提示型 bullet，只留真正的文案"""
    bullets = [str(b).strip() for b in (slide_data.get("content_bullets") or [])]
    real = [b for b in bullets if len(b) >= 6 and not META_HINT.search(b)]
    return real[:limit]


def fit_title_size(text: str, base: int = 40) -> int:
    """封面/致谢页标题按字数自动缩字号，避免长标题溢出文本框"""
    length = len(str(text or ""))
    if length <= 16:
        return base
    if length <= 24:
        return 32
    return 26


def fill_content_slide(slide, data: dict):
    for ph in slide.placeholders:
        idx = ph.placeholder_format.idx
        if idx == 0:
            ph.text = data.get("title", "")
        elif idx == 1:
            tf = ph.text_frame
            tf.clear()
            for i, bullet in enumerate(data.get("content_bullets") or []):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = bullet
                p.level = 0
    notes = data.get("notes")
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


async def render_with_progress(state: AgentState) -> dict:
    ppt_data = state["ppt_structure"]
    outline = ppt_data.get("slides") or []
    total = len(outline)
    if total == 0:
        raise ValueError("大纲为空，无法渲染 PPT")

    template_path = pick_template()
    template_name = os.path.basename(template_path)
    await update_progress(state, "RENDER", f"正在加载模板 {template_name} ...", sub_progress=30)

    prs = Presentation(template_path)
    xml_slides = prs.slides._sldIdLst
    orig_elems = list(xml_slides)
    n = len(orig_elems)

    # ---- 识别模板里的 封面 / 目录 / 致谢页 ----
    cover_idx = 0 if n >= 1 else None
    closing_idx = (n - 1) if n >= 2 else None
    toc_idx = None
    for i in range(1, max(n - 1, 1)):
        if "目录" in slide_text(prs.slides[i]):
            toc_idx = i
            break
    if toc_idx is not None and toc_idx == closing_idx:
        toc_idx = None

    # ---- 大纲页与模板页的映射：第1页→封面，最后1页→致谢，中间→正文 ----
    start = 1 if cover_idx is not None else 0
    if closing_idx is not None and (total - start) < 1:
        closing_idx = None
    end = total - 1 if closing_idx is not None else total
    if end < start:
        end = start
    middle = outline[start:end]
    use_toc = toc_idx is not None and len(middle) >= 2

    cover_data = outline[0] if cover_idx is not None else None
    closing_data = outline[-1] if closing_idx is not None else None

    # 保留要复用的模板页的 rId；示例页（Part 01~05 等）稍后删除
    keep_rids = set()
    if cover_idx is not None:
        keep_rids.add(orig_elems[cover_idx].rId)
    if use_toc:
        keep_rids.add(orig_elems[toc_idx].rId)
    if closing_idx is not None:
        keep_rids.add(orig_elems[closing_idx].rId)

    today = datetime.date.today().strftime("%Y年%m月%d日")
    reused = []

    # ---- 复用封面：加标题/副标题文本框（模板封面没有标题占位符）----
    if cover_idx is not None:
        cover = prs.slides[cover_idx]
        cover_title = cover_data.get("title", "")
        add_textbox(cover, cover_title, 0.6, 2.15, 12.1, 1.35,
                    fit_title_size(cover_title), bold=True, color=ACCENT)
        cover_sub = subtitle_lines(cover_data, 2)
        if cover_sub:
            add_textbox(cover, cover_sub, 0.6, 3.55, 12.1, 1.0, 18, color=SUBTITLE_COLOR)
        fill_date_chrome(cover, today)
        if cover_data.get("notes"):
            cover.notes_slide.notes_text_frame.text = cover_data["notes"]
        reused.append("封面")

    # ---- 复用目录页：条目替换成真实页标题 ----
    if use_toc:
        toc_slide = prs.slides[toc_idx]
        if fill_toc(toc_slide, [s.get("title", "") for s in middle]):
            reused.append("目录")
        else:
            # 找不到条目结构就不要这一页了
            keep_rids.discard(orig_elems[toc_idx].rId)
            use_toc = False

    # ---- 复用致谢页 ----
    if closing_idx is not None:
        closing_slide = prs.slides[closing_idx]
        closing_title = closing_data.get("title", "谢谢观看")
        add_textbox(closing_slide, closing_title, 0.6, 2.35, 12.1, 1.3,
                    fit_title_size(closing_title), bold=True, color=ACCENT)
        closing_sub = subtitle_lines(closing_data, 2)
        if closing_sub:
            add_textbox(closing_slide, closing_sub, 0.6, 3.7, 12.1, 0.9, 18, color=SUBTITLE_COLOR)
        fill_date_chrome(closing_slide, today)
        if closing_data.get("notes"):
            closing_slide.notes_slide.notes_text_frame.text = closing_data["notes"]
        reused.append("致谢页")

    # ---- 正文页：用模板版式新建 ----
    # 注意：必须在删除示例页【之前】新增。python-pptx 的新页 partname 是
    # len(sldIdLst)+1 算出来的，先删会让编号回退，撞上仍保留的 slideN.xml，
    # 导致 zip 里出现同名部件、致谢页内容被顶掉。
    for i, slide_data in enumerate(middle):
        page_num = start + i + 1
        sub_progress = 30 + int(((i + 1) / max(len(middle), 1)) * 60)
        await update_progress(
            state, "RENDER",
            f"正在渲染第 {page_num}/{total} 页：{slide_data.get('title')}",
            sub_progress=sub_progress,
        )
        layout = resolve_layout(prs, slide_data.get("layout_name", "标题和内容"))
        slide = prs.slides.add_slide(layout)
        fill_content_slide(slide, slide_data)

    added_elems = [e for e in list(xml_slides) if e.rId not in {o.rId for o in orig_elems}]

    # ---- 删除模板示例页 ----
    for elem in orig_elems:
        if elem.rId not in keep_rids:
            prs.part.drop_rel(elem.rId)
            xml_slides.remove(elem)

    # ---- 调整页序：封面 → 目录 → 正文 → 致谢 ----
    closing_elem = (orig_elems[closing_idx]
                    if closing_idx is not None and orig_elems[closing_idx].rId in keep_rids
                    else None)
    head = [orig_elems[cover_idx]] if (cover_idx is not None and orig_elems[cover_idx].rId in keep_rids) else []
    if use_toc:
        head.append(orig_elems[toc_idx])
    tail = [closing_elem] if closing_elem is not None else []
    for elem in list(xml_slides):
        xml_slides.remove(elem)
    for elem in head + added_elems + tail:
        xml_slides.append(elem)

    # ---- 兜底：按最终页序把所有 slide 部件重命名为连续的 slide1..N，杜绝同名部件 ----
    prs.part.rename_slide_parts([elem.rId for elem in xml_slides])

    # ---- 告诉前端这次套用了模板的哪些部分 ----
    await push_event(state, "template", {
        "template": template_name,
        "reused": reused,
        "content_pages": len(middle),
        "total_pages": len(list(xml_slides)),
    })

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"presentation_{state['session_id']}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
    output_path = os.path.join(OUTPUT_DIR, filename)

    await update_progress(state, "RENDER",
                          f"已套用模板 {template_name}（{'、'.join(reused) or '仅版式'}），正在保存文件...",
                          sub_progress=95)
    prs.save(output_path)

    progress_update = await update_progress(state, "DONE", "PPT生成完成！", sub_progress=100)
    log_step("RENDER", f"渲染完成: {output_path}", f"模板={template_name}, 复用={reused}")

    return {
        "output_path": output_path,
        "stage": "DONE",
        **progress_update,
    }
