# 目录: ./tests/inspect_template.py | 模块: 模板检查工具 | 职责: 输出模板主题字体、配色、封面形状布局，辅助渲染器开发
import zipfile
import re

from pptx import Presentation
from pptx.util import Emu

TEMPLATE = "templates/master.pptx"

z = zipfile.ZipFile(TEMPLATE)
theme_name = [n for n in z.namelist() if n.startswith("ppt/theme/")][0]
xml = z.read(theme_name).decode("utf-8")

mj = re.search(r"<a:majorFont>(.*?)</a:majorFont>", xml, re.S)
mn = re.search(r"<a:minorFont>(.*?)</a:minorFont>", xml, re.S)
for label, block in (("major", mj), ("minor", mn)):
    if not block:
        continue
    b = block.group(1)
    latin = re.search(r'typeface="([^"]*)"', b)
    ea = re.search(r'<a:ea typeface="([^"]*)"', b)
    print(f"{label} font -> latin={latin.group(1) if latin else None!r}, ea={ea.group(1) if ea else None!r}")

print("theme srgb colors:", re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml)[:12])

prs = Presentation(TEMPLATE)
print("slide size:", round(Emu(prs.slide_width).inches, 2), "x", round(Emu(prs.slide_height).inches, 2))

for idx in (0, 1, 7):
    slide = prs.slides[idx]
    print(f"=== slide{idx + 1} layout={slide.slide_layout.name}")
    for sh in slide.shapes:
        try:
            pos = (
                f"({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) "
                f"{Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f}in"
            )
        except Exception:
            pos = "n/a"
        txt = repr(sh.text_frame.text[:40]) if sh.has_text_frame else ""
        extra = ""
        if "AUTO_SHAPE" in str(sh.shape_type):
            try:
                extra = f" fill={sh.fill.fore_color.rgb}"
            except Exception as e:
                extra = f" fill=?({e})"
        print(f"   {sh.shape_type} {sh.name!r} {pos} {txt}{extra}")

cover_layout = prs.slides[0].slide_layout
print("cover layout shapes:", [(str(s.shape_type), s.name) for s in cover_layout.shapes])
