# 目录: ./tests/inspect_output.py | 模块: 产物检查 | 职责: 校验生成的 PPT 是否正确套用模板设计（封面/目录/致谢/正文）
import sys

from pptx import Presentation

path = sys.argv[1] if len(sys.argv) > 1 else "tests/last_output.pptx"
p = Presentation(path)

print(f"file={path} | slides={len(p.slides._sldIdLst)} | layouts={len(p.slide_layouts)}")
for i, s in enumerate(p.slides, 1):
    ph_texts = []
    for ph in s.placeholders:
        t = ph.text_frame.text.strip()
        if t:
            ph_texts.append(f"ph{ph.placeholder_format.idx}={t[:40]!r}")
    box_texts = [
        sh.text_frame.text.strip().replace("\n", " / ")[:60]
        for sh in s.shapes
        if sh.has_text_frame and not sh.is_placeholder and sh.text_frame.text.strip()
    ]
    notes = ""
    if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
        notes = f" | notes={s.notes_slide.notes_text_frame.text.strip()[:40]!r}"
    print(f"--- p{i} [{s.slide_layout.name}] placeholders: {ph_texts or 'None'}{notes}")
    if box_texts:
        print(f"      textboxes: {box_texts}")
