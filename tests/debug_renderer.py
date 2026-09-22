# 目录: ./tests/debug_renderer.py | 模块: 渲染调试 | 职责: 逐步打印 sldId/rId/partname，定位模板页复用与重排的底层问题
import os
import sys

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(BASE, "backend"))
os.chdir(os.path.join(BASE, "backend"))

from pptx import Presentation  # noqa: E402

from tools.ppt_renderer import remove_slides, resolve_layout, fill_content_slide  # noqa: E402


def dump(prs, label):
    xml = prs.slides._sldIdLst
    print(f"--- {label}")
    for elem in xml:
        rId = elem.rId
        try:
            part = prs.part.part_related_by(rId)
            pn = str(part.partname)
        except Exception as e:
            pn = f"<BROKEN: {e}>"
        print(f"    rId={rId} partname={pn}")
    # 包内所有 slide 部件
    pkg_partnames = sorted(
        str(p.partname) for p in prs.part.package.iter_parts()
        if str(p.partname).startswith("/ppt/slides/slide")
    )
    print(f"    package slide parts: {pkg_partnames}")


prs = Presentation("templates/master.pptx")
dump(prs, "原始模板")

orig_elems = list(prs.slides._sldIdLst)
cover_idx, toc_idx, closing_idx = 0, 2, 7
keep = [orig_elems[cover_idx], orig_elems[toc_idx], orig_elems[closing_idx]]
remove_slides(prs, keep)
dump(prs, "删除示例页后")

kept_before = list(prs.slides._sldIdLst)
print(f"kept_before count={len(kept_before)}")

titles = ["行业背景", "市场规模", "竞争格局", "技术趋势", "投资机会"]
new_slides = []
for t in titles:
    layout = resolve_layout(prs, "标题和内容")
    s = prs.slides.add_slide(layout)
    fill_content_slide(s, {"title": t, "content_bullets": ["a", "b"], "notes": "n"})
    new_slides.append(s)
dump(prs, "新增正文页后")

after = list(prs.slides._sldIdLst)
new_elems = after[len(kept_before):]
print(f"after count={len(after)} new_elems count={len(new_elems)}")

closing_elem = orig_elems[closing_idx]
print(f"closing_elem rId={closing_elem.rId} in xml: {closing_elem in list(prs.slides._sldIdLst)}")

xml = prs.slides._sldIdLst
head = [e for e in kept_before if e.rId != closing_elem.rId]
for elem in list(xml):
    xml.remove(elem)
for elem in head + new_elems + [closing_elem]:
    xml.append(elem)
dump(prs, "重排后")

os.makedirs("output", exist_ok=True)
prs.save("output/debug_reorder.pptx")
print("saved output/debug_reorder.pptx")

# 重新读回验证
p2 = Presentation("output/debug_reorder.pptx")
print("read back slides:", len(p2.slides._sldIdLst))
for i, s in enumerate(p2.slides, 1):
    print(f"  p{i} layout={s.slide_layout.name}")
