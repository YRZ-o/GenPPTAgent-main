# 目录: ./tests/inspect_toc_fonts.py | 模块: 模板字体检查 | 职责: 输出目录条目/标题框的字体规格，供渲染器自动适配
import re

from pptx import Presentation
from pptx.util import Emu, Pt

prs = Presentation("templates/master.pptx")


def dump(slide, label, names):
    print(f"=== {label}")
    for sh in slide.shapes:
        if not sh.has_text_frame or sh.name not in names:
            continue
        try:
            pos = f"({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f}in"
        except Exception:
            pos = "n/a"
        print(f"  {sh.name!r} {pos} align/hanchor text={sh.text_frame.text!r}")
        for pi, p in enumerate(sh.text_frame.paragraphs):
            for r in p.runs:
                f = r.font
                color = None
                try:
                    color = f.color.rgb
                except Exception:
                    color = getattr(f.color, "theme_color", None)
                ea = r._r.find(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr"
                )
                ea_type = None
                if ea is not None:
                    el = ea.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
                    if el is not None:
                        ea_type = el.get("typeface")
                print(f"    p{pi} run={r.text[:20]!r} name={f.name!r} ea={ea_type!r} "
                      f"size={f.size.pt if f.size else None} bold={f.bold} color={color}")


toc = prs.slides[1]
dump(toc, "目录页 (slide2)", {"目录：", "文本框 12", "文本框 11", "文本框 10", "文本框 9", "文本框 1"})

# 若为生成后的产物，额外打印所有条目框的几何与字号，检查溢出场景的克隆/重排/缩放
import sys  # noqa: E402

if len(sys.argv) > 1:
    out = Presentation(sys.argv[1])
    s = out.slides[1]
    print("=== 产物目录页条目框")
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        first = sh.text_frame.paragraphs[0].text if sh.text_frame.paragraphs else ""
        if re.match(r"^\s*(\d+)\s*[、.．]", first or "") or re.match(r"^\s*目录", first or ""):
            sizes = {r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size}
            colors = set()
            names = set()
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    try:
                        colors.add(str(r.font.color.rgb))
                    except Exception:
                        colors.add("?")
                    names.add(r.font.name)
            print(f"  {sh.name!r} top={Emu(sh.top).inches:.2f} text={sh.text_frame.text!r} "
                  f"sizes={sizes or 'inherit'} colors={colors} fonts={names}")

part = prs.slides[2]
dump(part, "Part 页 (slide3)", {"标题 1"})

cover = prs.slides[0]
dump(cover, "封面 (slide1)", {"文本框 4", "文本框 5"})
